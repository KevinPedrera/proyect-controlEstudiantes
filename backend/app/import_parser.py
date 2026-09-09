"""Parser acotado del perfil iDukay inspeccionado; sin persistencia ni guardado."""

from dataclasses import dataclass
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

PROFILE = "idukay_listado_filtrado_v1"
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5007
MAX_COLUMNS = 80
NAME_FIELDS = ("first_name", "middle_name", "last_name", "second_last_name")
STUDENT_FIELDS = ("document", "document_type", "last_name", "second_last_name", "first_name", "middle_name", "course", "parallel")
PERSON_FIELDS = ("last_name", "second_last_name", "first_name", "middle_name", "home_phone", "mobile_phone")
EMERGENCY_FIELDS = ("relationship", "first_name", "last_name", "mobile_phone", "home_phone", "work_phone")
PERSON_HEADERS = ("Apellido", "Segundo apellido", "Nombre", "Segundo nombre", "Teléfono casa", "Teléfono celular")
EMERGENCY_HEADERS = ("Parentesco", "Nombre", "Apellido", "Teléfono celular", "Teléfono domicilio", "Teléfono trabajo")
BLOCKS = (
    ("Información del estudiante", ("Cédula/Pasaporte", "Tipo de documento", "Apellido", "Segundo apellido", "Nombre", "Segundo nombre", "Curso", "Paralelo")),
    ("Información del representante legal", ("Parentesco",) + PERSON_HEADERS),
    ("Información del padre", PERSON_HEADERS),
    ("Información de la madre", PERSON_HEADERS),
    ("Información de contactos de emergencia", EMERGENCY_HEADERS * 2),
)


class ImportFormatError(ValueError):
    """Solo mensajes fijos/coordenadas; nunca incluir valores personales."""


def comparison_key(value: str | None) -> str:
    return " ".join((value or "").split()).casefold()


@dataclass
class ParsedImport:
    institution: str
    period: str
    period_key: str
    sheet: str
    rows: list[dict]
    skipped_empty_rows: int


def _check_archive(content: bytes) -> None:
    if len(content) > MAX_BYTES:
        raise ImportFormatError("El archivo supera el límite de 5 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(i.file_size for i in entries) > 40 * 1024 * 1024:
                raise ImportFormatError("El libro descomprimido excede los límites permitidos.")
            if any(i.flag_bits & 1 or 'vbaProject' in i.filename for i in entries):
                raise ImportFormatError("No se admiten libros cifrados o con macros.")
            for entry in entries:
                if entry.filename.endswith('.xml'):
                    raw = archive.read(entry)
                    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                        raise ImportFormatError("El libro contiene XML no admitido.")
                    if entry.filename.startswith('xl/worksheets/sheet'):
                        root = ElementTree.fromstring(raw)
                        ranges = [n.get('ref', '') for n in root.findall('.//{*}mergeCell') + root.findall('./{*}dimension')]
                        for area in ranges:
                            _, _, max_col, max_row = range_boundaries(area)
                            if max_col is None or max_row is None or max_col > MAX_COLUMNS or max_row > MAX_ROWS:
                                raise ImportFormatError("Las combinaciones o dimensiones exceden los límites permitidos.")
                        cells = root.findall('.//{*}c')
                        if len(cells) > MAX_ROWS * MAX_COLUMNS:
                            raise ImportFormatError("La hoja supera el límite de celdas.")
                        for cell in cells:
                            address = cell.get('r', '')
                            match = re.fullmatch(r'([A-Z]+)([0-9]+)', address)
                            if not match:
                                raise ImportFormatError("El libro contiene coordenadas incompatibles.")
                            column = 0
                            for char in match[1]:
                                column = column * 26 + ord(char) - 64
                            if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS:
                                raise ImportFormatError("La hoja supera 5007 filas u 80 columnas.")
    except ImportFormatError:
        raise
    except (BadZipFile, ElementTree.ParseError, KeyError, RuntimeError, EOFError, ValueError):
        raise ImportFormatError("No se pudo leer un archivo XLSX válido.") from None


def _layout(sheet, row: int, col: int) -> tuple[list[int], int]:
    merges = list(sheet.merged_cells.ranges)
    cursor = col
    columns = []
    bottom = row + 1
    for title, headers in BLOCKS:
        if comparison_key(sheet.cell(row, cursor).value) != comparison_key(title):
            raise ImportFormatError("Los bloques de encabezados no corresponden al perfil iDukay.")
        if not any(m.min_row == m.max_row == row and m.min_col == cursor and m.max_col == cursor + len(headers) - 1 for m in merges):
            raise ImportFormatError("Las combinaciones de grupos son incompatibles.")
        for offset, header in enumerate(headers):
            column = cursor + offset
            if comparison_key(sheet.cell(row + 1, column).value) != comparison_key(header):
                raise ImportFormatError("Falta un encabezado o su bloque es ambiguo.")
            vertical = [m for m in merges if m.min_col == m.max_col == column and m.min_row == row + 1]
            if len(vertical) != 1 or vertical[0].max_row != row + 2:
                raise ImportFormatError("Las combinaciones de columnas son incompatibles.")
            bottom = max(bottom, vertical[0].max_row)
            columns.append(column)
        cursor += len(headers)
    if comparison_key(sheet.cell(row + 1, col - 1).value) != '#':
        raise ImportFormatError("Falta el encabezado de numeración del listado.")
    return columns, bottom + 1


def _value(cell, warnings: list[dict]) -> str | None:
    value = cell.value
    if value is None or value == '':
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        warnings.append({"code": "NUMERIC_CELL", "cell": cell.coordinate, "message": "Celda numérica: revisar formato y posibles ceros perdidos."})
        value = str(int(value)) if value == int(value) else str(value)
    else:
        warnings.append({"code": "CELL_TYPE", "cell": cell.coordinate, "message": "Tipo de celda no admitido como dato; requiere revisión."})
        value = str(value)
    if len(value) > 240:
        raise ImportFormatError(f"El valor de {cell.coordinate} supera 240 caracteres.")
    if value == '-':
        warnings.append({"code": "LITERAL_DASH", "cell": cell.coordinate, "message": 'Guion literal "-" conservado como dato recibido.'})
    return value


def _parse_workbook(book) -> ParsedImport:
    if len(book.worksheets) > 10:
        raise ImportFormatError("El libro supera diez hojas.")
    candidates = []
    for sheet in book:
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
            raise ImportFormatError("Las dimensiones de la hoja exceden el perfil permitido.")
        for cells in sheet.iter_rows(min_row=1, max_row=min(40, sheet.max_row)):
            for cell in cells:
                if isinstance(cell.value, str) and comparison_key(cell.value) == comparison_key(BLOCKS[0][0]):
                    candidates.append((sheet, cell.row, cell.column))
    if len(candidates) != 1:
        raise ImportFormatError("Se requiere una única hoja y cabecera compatibles con iDukay.")
    sheet, header_row, start_column = candidates[0]
    columns, start_row = _layout(sheet, header_row, start_column)
    for cells in sheet:
        if any(c.data_type in ('f', 'e') for c in cells):
            raise ImportFormatError("El listado contiene fórmulas o errores Excel no soportados.")
    top = [c for cells in sheet.iter_rows(max_row=header_row - 1) for c in cells if isinstance(c.value, str) and c.value.strip()]
    titles = [c for c in top if comparison_key(c.value).startswith('listado de estudiantes')]
    periods = [(c, re.fullmatch(r'Año\s+Lectivo\s+(\d{4})\s*-\s*(\d{4})', c.value.strip(), re.I)) for c in top]
    periods = [(c, m) for c, m in periods if m]
    if len(titles) != 1 or len(periods) != 1:
        raise ImportFormatError("No se identifica un único título y año lectivo.")
    institution_cells = [c for c in top if c.column == titles[0].column and c.row < titles[0].row]
    if len(institution_cells) != 1:
        raise ImportFormatError("No se identifica inequívocamente la institución sobre el título.")
    institution = institution_cells[0].value.strip()
    if len(institution) > 240:
        raise ImportFormatError("El nombre institucional supera el límite permitido.")
    period_cell, match = periods[0]
    period_key = f'{match[1]}-{match[2]}'
    rows, skipped = [], 0
    for row_number in range(start_row, sheet.max_row + 1):
        cells = sheet[row_number]
        if all(c.value is None or isinstance(c.value, str) and not c.value.strip() for c in cells):
            skipped += 1
            continue
        if any(c.value is not None and str(c.value).strip() for c in cells if c.column < start_column - 1 or c.column > columns[-1]):
            raise ImportFormatError("Hay datos fuera de los bloques reconocidos.")
        warnings = []
        values = [_value(sheet.cell(row_number, col), warnings) for col in columns]
        student = dict(zip(STUDENT_FIELDS, values[:8]))
        required = {'first_name': 'nombre', 'last_name': 'apellido', 'course': 'curso', 'parallel': 'paralelo'}
        errors = [f'Falta {label}.' for field, label in required.items() if not student[field] or student[field] == '-']
        contacts = []
        offset = 8
        for kind, fields, slot in [('REPRESENTANTE_LEGAL', ('relationship',) + PERSON_FIELDS, None), ('PADRE', PERSON_FIELDS, None), ('MADRE', PERSON_FIELDS, None), ('EMERGENCIA', EMERGENCY_FIELDS, 1), ('EMERGENCIA', EMERGENCY_FIELDS, 2)]:
            contact = dict(zip(fields, values[offset:offset + len(fields)]))
            offset += len(fields)
            if any(v is not None for v in contact.values()):
                if not contact.get('first_name') or not contact.get('last_name'):
                    warnings.append({"code": "INCOMPLETE_CONTACT", "cell": f'{row_number}', "message": "Contacto con nombre o apellido sin registrar."})
                contacts.append({"type": kind, "source_slot": slot, **contact})
        rows.append({"row_number": row_number, "sheet": sheet.title, "student": student, "contacts": contacts, "warnings": warnings, "errors": errors})
    if not rows:
        raise ImportFormatError("El listado no contiene filas para previsualizar.")
    return ParsedImport(institution, period_cell.value.strip(), period_key, sheet.title, rows, skipped)


def parse_xlsx(content: bytes) -> ParsedImport:
    _check_archive(content)
    book = None
    try:
        book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        return _parse_workbook(book)
    except ImportFormatError:
        raise
    except Exception:
        raise ImportFormatError("No se pudo interpretar el perfil XLSX iDukay.") from None
    finally:
        if book is not None:
            book.close()
