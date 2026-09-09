"""Libros sintéticos: ningún dato del archivo institucional."""
from io import BytesIO

from openpyxl import Workbook

from app.import_parser import BLOCKS


def workbook(rows=None, shift=0):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Listado de estudiantes'
    for row, value, end in [(1, 'Colegio de prueba', 6), (2, 'Listado de estudiantes de Séptimo ( A )', 10), (3, 'Año Lectivo 2026 - 2027', 6)]:
        sheet.cell(row + shift, 2, value)
        sheet.merge_cells(start_row=row + shift, end_row=row + shift, start_column=2, end_column=end)
    sheet.cell(6 + shift, 2, '#')
    sheet.merge_cells(start_row=6 + shift, end_row=7 + shift, start_column=2, end_column=2)
    col = 3
    for title, headers in BLOCKS:
        sheet.cell(5 + shift, col, title)
        sheet.merge_cells(start_row=5 + shift, end_row=5 + shift, start_column=col, end_column=col + len(headers) - 1)
        for header in headers:
            sheet.cell(6 + shift, col, header)
            sheet.merge_cells(start_row=6 + shift, end_row=7 + shift, start_column=col, end_column=col)
            col += 1
    for offset, overrides in enumerate(rows if rows is not None else [{}]):
        data = {'B': f'{offset + 1}.', 'C': '0123456789', 'D': 'Cédula', 'E': 'Pérez', 'G': 'Ana', 'I': 'Séptimo', 'J': 'A', **overrides}
        for column, value in data.items():
            sheet[f'{column}{8 + shift + offset}'] = value
    return book


def binary(book):
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def upload(client, content=None, scope='SUBCONJUNTO'):
    return client.post('/api/student-imports/preview', files={'file': ('synthetic.xlsx', content if content is not None else binary(workbook()), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}, data={'scope': scope})
